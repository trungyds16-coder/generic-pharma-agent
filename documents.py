from __future__ import annotations

import html
import json
from pathlib import Path


def _iter_paragraphs(container):
    for paragraph in getattr(container, "paragraphs", []):
        yield paragraph
    for table in getattr(container, "tables", []):
        for row in table.rows:
            for cell in row.cells:
                yield from _iter_paragraphs(cell)


def _replace_in_paragraph(paragraph, values: dict[str, str]) -> None:
    full = "".join(run.text for run in paragraph.runs)
    updated = full
    for key, value in values.items():
        updated = updated.replace("{" + key + "}", value)
        updated = updated.replace("{{" + key + "}}", value)
    if updated != full and paragraph.runs:
        paragraph.runs[0].text = updated
        for run in paragraph.runs[1:]:
            run.text = ""


def fill_docx_template(template_path: str, output_path: str, fields_json: str,
                       structure_image_path: str = "") -> str:
    try:
        from docx import Document
        from docx.shared import Inches
    except ImportError as exc:
        raise RuntimeError("Thiếu python-docx; cài các phụ thuộc trong requirements.txt rồi thử lại.") from exc
    template = Path(template_path)
    output = Path(output_path)
    if not template.is_file():
        raise FileNotFoundError(f"Không tìm thấy template Word: {template}")
    try:
        fields = json.loads(fields_json or "{}")
    except json.JSONDecodeError as exc:
        raise ValueError(f"fields_json không hợp lệ: {exc}") from exc
    if not isinstance(fields, dict):
        raise ValueError("fields_json phải là một object JSON.")
    values = {str(k): str(v) for k, v in fields.items() if v is not None}
    doc = Document(str(template))
    containers = [doc]
    for section in doc.sections:
        containers.extend([section.header, section.footer])
    for container in containers:
        for paragraph in _iter_paragraphs(container):
            _replace_in_paragraph(paragraph, values)
    image_path = Path(structure_image_path) if structure_image_path else None
    if image_path:
        if not image_path.is_file():
            raise FileNotFoundError(f"Không tìm thấy ảnh cấu trúc: {image_path}")
        paragraphs = list(_iter_paragraphs(doc))
        for index, paragraph in enumerate(paragraphs):
            if "Cấu trúc hoạt chất sử dụng" in paragraph.text and index + 1 < len(paragraphs):
                target = paragraphs[index + 1]
                target.text = ""
                target.add_run().add_picture(str(image_path), width=Inches(2.5))
                break
        else:
            raise ValueError("Không tìm thấy vị trí 'Cấu trúc hoạt chất sử dụng' trong template.")
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    return f"Đã tạo Word: {output}; đã điền {len(values)} trường" + (" và chèn ảnh cấu trúc." if image_path else ".")


def fill_html_template(template_path: str, output_path: str, fields_json: str) -> str:
    template = Path(template_path)
    if not template.is_file():
        raise FileNotFoundError(f"Không tìm thấy template HTML: {template}")
    fields = json.loads(fields_json or "{}")
    source = template.read_text(encoding="utf-8", errors="replace")
    original = source
    for key, raw in fields.items():
        value = html.escape(str(raw if raw is not None else ""))
        source = source.replace("{{" + str(key) + "}}", value)
        source = source.replace("{" + str(key) + "}", value)
    if source == original:
        raise ValueError("Template HTML không có placeholder tương ứng; dùng create_html_report để tạo báo cáo mới từ dữ liệu.")
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(source, encoding="utf-8")
    return f"Đã tạo HTML theo template: {output}"


def convert_pdf_to_html(pdf_path: str, output_path: str, title: str = "Tài liệu PDF") -> str:
    try:
        import fitz
    except ImportError as exc:
        raise RuntimeError("Thiếu PyMuPDF; cài các phụ thuộc trong requirements.txt rồi thử lại.") from exc
    source = Path(pdf_path)
    output = Path(output_path)
    if not source.is_file() or source.suffix.lower() != ".pdf":
        raise FileNotFoundError(f"Không tìm thấy file PDF: {source}")
    output.parent.mkdir(parents=True, exist_ok=True)
    safe_title = html.escape(title or source.stem)
    with fitz.open(source) as pdf, output.open("w", encoding="utf-8", newline="") as target:
        page_count = len(pdf)
        nav = " · ".join(f'<a href="#page-{n}">{n}</a>' for n in range(1, page_count + 1))
        target.write(
            '<!doctype html><html lang="vi"><head><meta charset="utf-8">'
            '<meta name="viewport" content="width=device-width,initial-scale=1">'
            f'<title>{safe_title}</title><style>'
            'body{font:16px/1.65 system-ui,sans-serif;max-width:1000px;margin:2rem auto;padding:0 1rem;color:#18212f}'
            'nav{position:sticky;top:0;background:#fff;padding:.8rem;border-bottom:1px solid #ddd}'
            '.page{border-bottom:1px solid #ddd;padding:1rem 0;scroll-margin-top:4rem}'
            'pre{white-space:pre-wrap;overflow-wrap:anywhere;font:inherit}h2{color:#1b5c88}'
            f'</style></head><body><h1>{safe_title}</h1><p>Nguồn: {html.escape(source.name)} · '
            f'{page_count} trang · Có thể tìm nội dung bằng Ctrl+F.</p><nav>{nav}</nav>'
        )
        for number, page in enumerate(pdf, start=1):
            text = page.get_text("text", sort=True).strip()
            escaped = html.escape(text) if text else "(Trang này không có lớp văn bản; có thể là PDF scan.)"
            target.write(f'<section class="page" id="page-{number}"><h2>Trang {number}</h2><pre>{escaped}</pre></section>')
        target.write('</body></html>')
    return f"Đã chuyển {page_count} trang sang HTML có thể tìm kiếm: {output}. Trang scan được đánh dấu để nhận diện OCR riêng."


def create_html_report(output_path: str, title: str, sections_json: str, sources_json: str = "", image_path: str = "") -> str:
    try:
        sections = json.loads(sections_json or "{}")
        sources = json.loads(sources_json or "[]")
    except json.JSONDecodeError as exc:
        raise ValueError(f"JSON báo cáo không hợp lệ: {exc}") from exc
    if not isinstance(sections, dict):
        raise ValueError("sections_json phải là object JSON theo từng mục báo cáo.")
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    safe_title = html.escape(title or "Báo cáo")
    blocks=[]
    for heading, value in sections.items():
        if isinstance(value, list):
            content = "<ul>" + "".join(f"<li>{html.escape(str(item))}</li>" for item in value) + "</ul>"
        elif isinstance(value, dict):
            content = "<dl>" + "".join(f"<dt>{html.escape(str(k))}</dt><dd>{html.escape(str(v))}</dd>" for k,v in value.items()) + "</dl>"
        else:
            content = "<p>" + html.escape(str(value)).replace("\n", "<br>") + "</p>"
        blocks.append(f"<section><h2>{html.escape(str(heading))}</h2>{content}</section>")
    image_block=""
    if image_path:
        image=Path(image_path)
        if not image.is_file(): raise FileNotFoundError(f"Không tìm thấy ảnh: {image}")
        try: image_ref=image.resolve().relative_to(out.parent.resolve()).as_posix()
        except ValueError: raise ValueError("Ảnh phải nằm trong project cùng báo cáo để HTML tự mở được.")
        image_block=f'<figure><img src="{html.escape(image_ref)}" alt="Công thức cấu tạo phân tử"><figcaption>Công thức cấu tạo · xem nguồn ở danh mục tài liệu</figcaption></figure>'
    source_items=[]
    if isinstance(sources,list):
        for item in sources:
            if isinstance(item,dict): url=str(item.get("url", "")); label=str(item.get("title") or url); date=str(item.get("date") or item.get("recorded_at", ""))
            else: url=str(item); label=url; date=""
            if url.startswith(("https://","http://")):
                source_items.append(f'<li><a href="{html.escape(url,quote=True)}" rel="noopener noreferrer">{html.escape(label)}</a>{(" · " + html.escape(date)) if date else ""}</li>')
    sources_block="<section><h2>Tài liệu tham khảo</h2><ol>"+"".join(source_items)+"</ol></section>" if source_items else ""
    result=(
        '<!doctype html><html lang="vi"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">'
        f'<title>{safe_title}</title><style>body{{font:16px/1.7 system-ui,sans-serif;max-width:980px;margin:2rem auto;padding:0 1.2rem;color:#1e293b}}h1{{color:#123d5a;border-bottom:3px solid #2a7999;padding-bottom:.6rem}}h2{{color:#155e75;margin-bottom:.4rem}}section{{margin:1.6rem 0;padding:0 0 1rem;border-bottom:1px solid #dbe3e8}}figure{{text-align:center}}figure img{{max-width:420px;max-height:360px}}dt{{font-weight:700}}dd{{margin-bottom:.6rem}}a{{color:#075985}}</style></head><body>'
        f'<h1>{safe_title}</h1>{image_block}{"".join(blocks)}{sources_block}<footer>Được tạo cục bộ · Hãy đối chiếu dữ liệu chuyên môn với nguồn gốc.</footer></body></html>'
    )
    out.write_text(result,encoding="utf-8")
    return f"Đã tạo báo cáo HTML: {out}; {len(blocks)} mục, {len(source_items)} nguồn."


def read_document(path: str, max_chars: int = 30000) -> str:
    source=Path(path)
    if not source.is_file(): raise FileNotFoundError(f"Không tìm thấy tài liệu: {source}")
    suffix=source.suffix.lower()
    if suffix==".docx":
        try: from docx import Document
        except ImportError as exc: raise RuntimeError("Thiếu python-docx.") from exc
        doc=Document(str(source)); parts=[]
        for paragraph in _iter_paragraphs(doc):
            if paragraph.text.strip(): parts.append(paragraph.text.strip())
        text="\n".join(parts)
    elif suffix==".pdf":
        try: import fitz
        except ImportError as exc: raise RuntimeError("Thiếu PyMuPDF.") from exc
        with fitz.open(source) as pdf:
            text="\n\n".join(f"[Trang {i+1}]\n"+page.get_text("text",sort=True) for i,page in enumerate(pdf))
    elif suffix in {".html",".htm"}:
        from research import _TextParser
        parser=_TextParser(); parser.feed(source.read_text(encoding="utf-8",errors="replace")); text=" ".join(parser.parts)
    else:
        text=source.read_text(encoding="utf-8",errors="replace")
    return f"Tài liệu: {source}\n\n{text[:max_chars]}"+("\n...[đã rút gọn]" if len(text)>max_chars else "")
